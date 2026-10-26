"""Config, CSV parsing, outcome labels, lead-card formatting, keyboards, post_next_lead.

Shared by admin_bot.py, crm_bot.py and dashboard.py. Plain text only.
"""

import csv
import io
import json
import os
import re
from datetime import datetime, timezone

import openpyxl
import phonenumbers
from dotenv import load_dotenv
from phonenumbers import timezone as phone_tz
from telegram import Bot, InlineKeyboardButton, InlineKeyboardMarkup

import db

load_dotenv()
ADMIN_BOT_TOKEN = os.environ["ADMIN_BOT_TOKEN"]
CRM_BOT_TOKEN = os.environ["CRM_BOT_TOKEN"]
ADMIN_USER_ID = int(os.environ["ADMIN_USER_ID"])
TURSO_DATABASE_URL = os.environ["TURSO_DATABASE_URL"]
TURSO_AUTH_TOKEN = os.environ["TURSO_AUTH_TOKEN"]
DASHBOARD_KEY = os.environ["DASHBOARD_KEY"]

# Google Voice texts arrive through a mautrix-gvoice bridge on a Matrix
# homeserver (deploy/gvoice). All optional: with no homeserver set, gvoice.py
# does not start and the dashboard's Texts page says so.
MATRIX_HOMESERVER = (os.environ.get("MATRIX_HOMESERVER") or "").strip().rstrip("/")
MATRIX_USER = (os.environ.get("MATRIX_USER") or "crm").strip()
MATRIX_PASSWORD = os.environ.get("MATRIX_PASSWORD") or ""
# The bridge's ghost users look like @gvoice_<login>.+15551234567:<domain> and
# the bridge bot is @gvoicebot:<domain>, from appservice.username_template and
# appservice.bot.username in the bridge config.
GVOICE_GHOST_PREFIX = (os.environ.get("GVOICE_GHOST_PREFIX") or "gvoice_").strip()
GVOICE_BOT = (os.environ.get("GVOICE_BOT") or "gvoicebot").strip().lstrip("@")
GVOICE_SOURCE = (os.environ.get("GVOICE_SOURCE") or "google-voice").strip() or "google-voice"
# Post every inbound text to the allowed group chats so callers see replies.
GVOICE_ANNOUNCE = (os.environ.get("GVOICE_ANNOUNCE") or "1").strip().lower() not in ("0", "no", "false", "")

OUTCOMES = {  # callback key -> button label
    "no_response": "No R",
    "hung_up": "Hung up",
    "on_phone": "On the phone",
}


# --- upload parsing (any file shape) ------------------------------------------

_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_DATE_RE = re.compile(r"^\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}")
_PHONE_HEADER_WORDS = db.PHONE_HEADER_WORDS

# The country an upload's local-looking numbers are read as, unless the upload
# form says otherwise. A value that arrives with + or 00 carries its own
# country code and is parsed with that instead.
DEFAULT_COUNTRY = (os.environ.get("DEFAULT_COUNTRY") or "GB").strip().upper()

# What the upload form accepts. Spreadsheets, delimited text and JSON all end
# up as the same list of row dicts, so the rest of the app sees one shape.
UPLOAD_EXTENSIONS = (
    ".csv", ".tsv", ".tab", ".txt", ".xlsx", ".xlsm", ".json", ".jsonl", ".ndjson",
)

# A row is written to the database only when its status is "valid".
ROW_STATUSES = ("valid", "invalid", "wrong_country", "duplicate", "blocked", "no_phone")


def _decode(raw) -> str:
    """Bytes -> text, UTF-8 with any Excel byte-order mark gone, else latin-1."""
    if not isinstance(raw, (bytes, bytearray)):
        return raw or ""
    try:
        return bytes(raw).decode("utf-8-sig")
    except UnicodeDecodeError:
        return bytes(raw).decode("latin-1")


def _cell_text(value) -> str:
    """One cell from any source -> trimmed text.

    Spreadsheets and JSON hand over real numbers, so 448000000000.0 has to lose
    its ".0" before it can look like a phone number, and dates arrive as
    datetimes rather than the strings a CSV would carry.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def _digits_and_plus(value: str) -> str:
    """Keep a leading + and digits only. Also unwraps Excel's 4.48E+11 form."""
    value = _cell_text(value)
    if re.fullmatch(r"\d+(\.\d+)?[eE]\+?\d+", value):
        try:
            value = str(int(float(value)))
        except (ValueError, OverflowError):
            pass
    plus = value.startswith("+")
    digits = re.sub(r"\D", "", value)
    return ("+" if plus else "") + digits


def _looks_like_phone(value: str) -> bool:
    value = _cell_text(value)
    if _DATE_RE.match(value) or _EMAIL_RE.match(value):
        return False
    d = _digits_and_plus(value)
    return 8 <= len(d.lstrip("+")) <= 15 and d.lstrip("+").isdigit()


def _first_row_is_header(cells: list) -> bool:
    """A row holding an email, a phone number or a date is data, not a header."""
    for cell in cells:
        cell = _cell_text(cell)
        if _EMAIL_RE.match(cell) or _DATE_RE.match(cell) or _looks_like_phone(cell):
            return False
    return True


def _infer_headers(cells: list) -> list[str]:
    """Name columns from the shape of the first data row."""
    names, seen = [], {}
    for i, cell in enumerate(cells, start=1):
        cell = _cell_text(cell)
        if _EMAIL_RE.match(cell):
            name = "Email"
        elif _looks_like_phone(cell):
            name = "Phone"
        elif _DATE_RE.match(cell):
            name = "Date"
        elif re.fullmatch(r"[A-Z]{2}", cell):
            name = "Region"
        else:
            name = f"Column {i}"
        seen[name] = seen.get(name, 0) + 1
        names.append(name if seen[name] == 1 else f"{name} {seen[name]}")
    return names


def _clean_row(row: dict) -> dict:
    """Trim a row dict's keys and values, dropping keys that trim to nothing."""
    clean = {}
    for key, value in row.items():
        key = _cell_text(key)
        if not key:
            continue
        clean[key] = _cell_text(value)
    return clean


def _clean_cells(cells: list, headers: list[str]) -> dict:
    """One row of cells + the header names -> a row dict.

    Cells past the last header are kept as "Column N", matching what the CSV
    reader used to do with them.
    """
    clean = {}
    for i, cell in enumerate(cells):
        name = headers[i] if i < len(headers) else f"Column {i + 1}"
        value = _cell_text(cell)
        if i >= len(headers) and not value:
            continue
        if name:
            clean[name] = value
    return clean


def _rows_from_matrix(matrix: list[list]) -> list[dict]:
    """A grid of cells -> row dicts, working out whether row one is a header."""
    if not matrix:
        return []
    first = matrix[0]
    if _first_row_is_header(first):
        headers = [_cell_text(c) for c in first]
        body = matrix[1:]
    else:
        # No header row: name the columns from the data and keep the first row.
        headers = _infer_headers(first)
        body = matrix
    rows = [_clean_cells(cells, headers) for cells in body]
    return [r for r in rows if any(r.values())]


def _sniff_delimiter(text: str) -> str:
    """Comma, tab, semicolon or pipe, from the first few lines."""
    sample = "\n".join(text.splitlines()[:20])
    try:
        return csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
    except csv.Error:
        # One column per line, or something the sniffer cannot read: a delimiter
        # that does not appear leaves every line as a single value.
        return ","


def parse_delimited(text: str) -> list[dict]:
    """CSV, TSV or one-value-per-line text -> row dicts."""
    text = _decode(text).lstrip("\ufeff")
    if not text.strip():
        return []
    reader = csv.reader(io.StringIO(text), delimiter=_sniff_delimiter(text))
    return _rows_from_matrix([row for row in reader if row])


def parse_sheet(raw: bytes) -> list[dict]:
    """The first sheet of an .xlsx/.xlsm workbook -> row dicts."""
    book = openpyxl.load_workbook(
        io.BytesIO(bytes(raw)), read_only=True, data_only=True
    )
    try:
        sheet = book[book.sheetnames[0]]
        matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        book.close()
    # read_only sheets pad every row out to the widest one with Nones.
    matrix = [row for row in matrix if any(_cell_text(c) for c in row)]
    return _rows_from_matrix(matrix)


def parse_json(text: str) -> list[dict]:
    """A JSON array of objects, or one JSON object per line -> row dicts.

    Objects keep their own keys, so the file brings its own schema. An array of
    arrays is read as a grid, the same as a CSV.
    """
    text = _decode(text).strip()
    if not text:
        return []
    try:
        loaded = json.loads(text)
    except ValueError:
        loaded = []
        for line in text.splitlines():
            line = line.strip().rstrip(",")
            if not line:
                continue
            try:
                loaded.append(json.loads(line))
            except ValueError as exc:
                raise ValueError(f"Line {len(loaded) + 1} is not valid JSON: {exc}") from exc
    if isinstance(loaded, dict):
        # A single object, or a wrapper with the list under some key.
        for value in loaded.values():
            if isinstance(value, list) and value:
                loaded = value
                break
        else:
            loaded = [loaded]
    if not isinstance(loaded, list):
        raise ValueError("That JSON is not a list of records.")
    if loaded and all(isinstance(item, (list, tuple)) for item in loaded):
        return _rows_from_matrix([list(item) for item in loaded])
    rows = []
    for item in loaded:
        if isinstance(item, dict):
            rows.append(_clean_row(item))
        elif item is not None:
            rows.append({"Column 1": _cell_text(item)})
    return [r for r in rows if any(r.values())]


def parse_upload(filename: str, raw) -> list[dict]:
    """Any supported file -> row dicts, picked by extension then by content.

    Every format lands on the same shape, so a file brings whatever columns it
    likes and the rest of the app does not care which format it came from.
    """
    ext = os.path.splitext(filename or "")[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return parse_sheet(raw)
    if ext in (".json", ".jsonl", ".ndjson"):
        return parse_json(raw)
    if ext in (".csv", ".tsv", ".tab", ".txt"):
        return parse_delimited(raw)
    # No extension we know: a zip header means a workbook, a leading [ or {
    # means JSON, anything else is read as delimited text.
    head = bytes(raw)[:4] if isinstance(raw, (bytes, bytearray)) else (raw or "")[:4].encode()
    if head[:2] == b"PK":
        return parse_sheet(raw)
    if head[:1] in (b"[", b"{"):
        return parse_json(raw)
    return parse_delimited(raw)


def parse_csv(text) -> list[dict]:
    """Kept so the admin bot and anything else written against v1 still work."""
    return parse_delimited(text)


# --- phone numbers and row checks ---------------------------------------------


def to_e164(value, country: str) -> tuple[str, str, tuple[str, ...]]:
    """(E.164 number, the ISO country it belongs to, its time zones).

    ("", "", ()) when the value is not a real, dialable number. `country` is the
    region local-looking numbers are read as; a value carrying + or 00 already
    names its own country code and is parsed with that instead.
    """
    raw = _digits_and_plus(value)
    if not raw.lstrip("+"):
        return "", "", ()
    if raw.startswith("00"):
        raw = "+" + raw[2:]
    region = None if raw.startswith("+") else (country or None)
    try:
        num = phonenumbers.parse(raw, region)
    except phonenumbers.NumberParseException:
        return "", "", ()
    if not phonenumbers.is_valid_number(num):
        return "", "", ()
    zones = tuple(z for z in phone_tz.time_zones_for_number(num) if z != "Etc/Unknown")
    return (
        phonenumbers.format_number(num, phonenumbers.PhoneNumberFormat.E164),
        phonenumbers.region_code_for_number(num) or "",
        zones,
    )


def country_timezones(country: str) -> tuple[str, ...]:
    """The time zones a country uses, for the upload form's default."""
    try:
        code = phonenumbers.country_code_for_region((country or "").upper())
    except Exception:
        return ()
    if not code:
        return ()
    example = phonenumbers.example_number((country or "").upper())
    if example is None:
        return ()
    return tuple(z for z in phone_tz.time_zones_for_number(example) if z != "Etc/Unknown")


def _dial_code(region: str) -> int:
    """The calling code of an ISO country, 0 when there is not one."""
    try:
        return phonenumbers.country_code_for_region((region or "").upper()) or 0
    except Exception:
        return 0


def _shares_dial_code(region: str, country: str) -> bool:
    """True when two ISO countries share a calling code.

    +44 covers GB, GG, JE and IM, so a real UK mobile can report its region as
    Guernsey. Comparing the calling code instead of the region keeps those
    numbers out of the "wrong country" pile.
    """
    if not region or not country:
        return False
    if region.upper() == country.upper():
        return True
    try:
        return phonenumbers.country_code_for_region(
            region.upper()
        ) == phonenumbers.country_code_for_region(country.upper())
    except Exception:
        return False


def classify_row(checks: dict) -> tuple[str, str]:
    """Decide one row's status and the reason the preview shows for it.

    `checks` describes one row:
      has_digits      bool - the phone cell held any digits at all
      e164            str  - the number in E.164 form, or "" when it is not a
                             real dialable number for the chosen country
      region          str  - the ISO country the number actually belongs to,
                             "" when it could not be worked out
      country         str  - the ISO country chosen on the upload form
      same_dial_code  bool - region and country share a calling code, so a +44
                             Guernsey-flagged UK mobile counts as matching
      strict          bool - the form's "only numbers from this country" toggle
      seen_in_file    bool - an earlier row in this same file had the same key
      blocked         bool - the number is on the do-not-even-try list

    Return a (status, reason) pair. `status` must be one of ROW_STATUSES, and
    only "valid" rows are written to the database. `reason` is shown to the user
    as written, so phrase it for a person: "not a real AU number", "already
    flagged", "same number as an earlier row".

    Several checks can be true at once, so the order below is the policy. A
    block is a decision somebody made about that number, so it is reported ahead
    of the file merely repeating itself; otherwise a flagged number arriving
    twice would be written off as a duplicate and the flag never seen. The
    country check comes last of the held-back reasons because it is the only one
    the uploader can undo from the form itself.
    """
    if not checks.get("has_digits"):
        return "no_phone", "no number in that column"
    if not checks.get("e164"):
        return "invalid", "not a real, dialable number"
    if checks.get("blocked"):
        return "blocked", "already flagged do not even try"
    if checks.get("seen_in_file"):
        return "duplicate", "same number as an earlier row"
    if checks.get("strict") and not checks.get("same_dial_code"):
        country = checks.get("country") or ""
        code = checks.get("dial_code") or 0
        own = checks.get("country_dial_code") or 0
        if code:
            return "wrong_country", (
                f"a +{code} number, {country} is +{own}" if own
                else f"a +{code} number, not {country}"
            )
        return "wrong_country", "from another country"
    return "valid", ""


def validate_rows(
    rows: list[dict],
    country: str = "",
    strict: bool = False,
    phone_column: str = "",
    blocked=(),
) -> list[dict]:
    """Check every row before anything is written, and say what is wrong.

    Returns one result dict per row: the row itself (its phone cell rewritten to
    E.164 when the number is real), the digits-only `key` the database leads are
    keyed on, the number's own country and time zones, and the status and reason
    from classify_row. Nothing here touches the database.
    """
    country = (country or DEFAULT_COUNTRY).strip().upper()
    rows = [dict(r) for r in rows if r]
    if not rows:
        return []
    column = db.pick_phone_column(rows, phone_column)
    blocked = set(blocked)
    results, seen = [], set()
    for row in rows:
        raw = row.get(column) or ""
        e164, region, zones = to_e164(raw, country)
        key = db.normalize_phone(e164 or raw)
        status, reason = classify_row(
            {
                "has_digits": bool(_digits_and_plus(raw).lstrip("+")),
                "e164": e164,
                "region": region,
                "dial_code": _dial_code(region),
                "country": country,
                "country_dial_code": _dial_code(country),
                "same_dial_code": _shares_dial_code(region, country),
                "strict": bool(strict),
                "seen_in_file": bool(key) and key in seen,
                "blocked": bool(key) and key in blocked,
            }
        )
        if key:
            seen.add(key)
        if e164:
            row = dict(row, **{column: e164})
        results.append(
            {
                "row": row,
                "phone": e164,
                "key": key,
                "region": region,
                "dial_code": _dial_code(region),
                "timezones": zones,
                "status": status,
                "reason": reason,
                "raw": _cell_text(raw),
            }
        )
    return results


def count_statuses(results: list[dict]) -> dict:
    """How many rows landed in each status, for the preview's summary line."""
    counts = {s: 0 for s in ROW_STATUSES}
    for r in results:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    return counts


# --- formatting -----------------------------------------------------------------


def fmt_talk(seconds: int | None) -> str:
    """Seconds -> 'm:ss', or 'not answered' when None."""
    if seconds is None:
        return "not answered"
    seconds = max(0, int(seconds))
    return f"{seconds // 60}:{seconds % 60:02d}"


def _local_hhmm(utc_text: str) -> str:
    """'YYYY-MM-DD HH:MM:SS' (UTC) -> 'HH:MM' in the server's local time."""
    try:
        dt = datetime.strptime(utc_text, "%Y-%m-%d %H:%M:%S").replace(tzinfo=timezone.utc)
        return dt.astimezone().strftime("%H:%M")
    except (TypeError, ValueError):
        return str(utc_text)[11:16]


def _history_line(past: dict) -> str:
    date = str(past.get("ended_at") or past.get("claimed_at") or "")[:10]
    caller = past.get("caller_name") or str(past.get("caller_id") or "")
    outcome = past.get("outcome")
    label = OUTCOMES.get(outcome, outcome) if outcome else "no outcome"
    talk = fmt_talk(past.get("talk_seconds"))
    notes = " ".join(str(past.get("notes") or "").split())[:80]
    return f"{date}  {caller}  {label}  {talk}  {notes}".rstrip()


def format_lead(lead: dict, attempt: dict | None = None, history: list[dict] | None = None) -> str:
    """The lead card.

    Lead #<id>  [source: <name>], the CSV fields one per line, a 'Seen before:'
    block when history is non-empty, then the attempt's state: 'Claimed by <name>',
    'On the line since <HH:MM>' while answered, or 'Outcome: <label>  Talk: <m:ss>'
    once ended ('Outcome: <label>  not answered' when talk_seconds is NULL).
    """
    source = lead.get("source") or ""
    head = f"Lead #{lead['id']}"
    if source:
        head += f"  [source: {source}]"
    lines = [head]
    for header, value in (lead.get("data") or {}).items():
        if value is None:
            continue
        value = str(value).strip()
        if not value:
            continue
        lines.append(f"{header}: {value}")

    if history:
        lines.append("")
        lines.append("Seen before:")
        for past in history:
            lines.append(_history_line(past))

    if attempt:
        lines.append("")
        lines.append(f"Claimed by {attempt.get('caller_name') or attempt.get('caller_id')}")
        if attempt.get("ended_at"):
            outcome = attempt.get("outcome")
            label = OUTCOMES.get(outcome, outcome) if outcome else "no outcome"
            if attempt.get("talk_seconds") is None:
                lines.append(f"Outcome: {label}  not answered")
            else:
                lines.append(f"Outcome: {label}  Talk: {fmt_talk(attempt['talk_seconds'])}")
        elif attempt.get("answered_at"):
            lines.append(f"On the line since {_local_hhmm(attempt['answered_at'])}")
    return "\n".join(lines)


# --- keyboards ------------------------------------------------------------------


def posted_keyboard(lead_id: int) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        [
            [
                InlineKeyboardButton("I'm calling it", callback_data=f"claim:{lead_id}"),
                InlineKeyboardButton("Do not even try", callback_data=f"flag:{lead_id}"),
            ]
        ]
    )


def claimed_keyboard(attempt_id: int, lead_id: int) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        [
            [
                InlineKeyboardButton("Answered", callback_data=f"answered:{attempt_id}"),
                InlineKeyboardButton(
                    OUTCOMES["no_response"], callback_data=f"outcome:{attempt_id}:no_response"
                ),
                InlineKeyboardButton("Do not even try", callback_data=f"flag:{lead_id}"),
            ]
        ]
    )


def answered_keyboard(attempt_id: int) -> InlineKeyboardMarkup:
    row = [
        InlineKeyboardButton(OUTCOMES[key], callback_data=f"outcome:{attempt_id}:{key}")
        for key in ("hung_up", "no_response", "on_phone")
    ]
    return InlineKeyboardMarkup([row])


# --- posting --------------------------------------------------------------------


async def post_next_lead(chat_id: int) -> bool:
    """Post the next new lead (with its history) to chat_id; False when none is left."""
    lead = db.next_new_lead()
    if lead is None:
        return False
    history = db.history_for_phone(lead["phone"])
    sent = await Bot(CRM_BOT_TOKEN).send_message(
        chat_id,
        format_lead(lead, None, history),
        reply_markup=posted_keyboard(lead["id"]),
    )
    db.mark_posted(lead["id"], chat_id, sent.message_id)
    return True
